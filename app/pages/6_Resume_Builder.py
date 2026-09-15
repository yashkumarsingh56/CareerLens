import streamlit as st
import io
from docx import Document
from docx.shared import Pt

st.set_page_config(page_title="Resume Builder", page_icon="📝", layout="wide")

st.title("📝 CareerLens Resume Builder")
st.markdown("Build a clean, ATS-friendly resume tailored to your target IT role. Fill in your details below and download your formatted resume.")

# Default Skills Pool (matching our analytics)
SKILL_VOCAB = [
    'SQL', 'Python', 'Java', 'JavaScript', 'React', 'Node.js', 'MongoDB', 
    'Express.js', 'Power BI', 'Tableau', 'Excel', 'AWS', 'Azure', 'GCP', 
    'Docker', 'Kubernetes', 'C++', 'C#', 'Statistics', 'Machine Learning', 
    'TensorFlow', 'PyTorch'
]

col1, col2 = st.columns([1, 1])

with col1:
    st.header("1. Personal Details")
    full_name = st.text_input("Full Name", "John Doe")
    email = st.text_input("Email", "john.doe@example.com")
    phone = st.text_input("Phone Number", "+1 234 567 8900")
    linkedin = st.text_input("LinkedIn URL", "linkedin.com/in/johndoe")
    github = st.text_input("GitHub URL", "github.com/johndoe")

    st.header("2. Professional Summary")
    summary = st.text_area("Brief summary of your career and goals", 
                           "Results-driven software engineer with 3+ years of experience building scalable web applications...")

    st.header("3. Skills")
    selected_skills = st.multiselect("Select your technical skills", SKILL_VOCAB, default=['Python', 'SQL', 'AWS'])
    other_skills = st.text_input("Other Skills (comma separated)", "Git, Agile, REST APIs")

with col2:
    st.header("4. Experience")
    st.markdown("Add your most recent role.")
    job_title = st.text_input("Job Title", "Software Engineer")
    company = st.text_input("Company", "Tech Innovations Inc.")
    job_dates = st.text_input("Dates (e.g., Jan 2021 - Present)", "Jan 2021 - Present")
    job_desc = st.text_area("Responsibilities & Achievements (bullet points)", 
                            "- Developed scalable backend APIs using Node.js and Express\n- Improved database query performance by 40% using optimized SQL indexes\n- Deployed applications using Docker and AWS CI/CD pipelines.")
    
    st.header("5. Education")
    degree = st.text_input("Degree", "B.S. Computer Science")
    university = st.text_input("University", "State University")
    edu_dates = st.text_input("Graduation Date", "May 2020")

st.markdown("---")

# Generate Markdown Resume
st.header("Preview & Download")

# Combine skills
all_skills = selected_skills.copy()
if other_skills.strip():
    all_skills.extend([s.strip() for s in other_skills.split(',')])
skills_str = " | ".join(all_skills)

resume_md = f"""
# {full_name}
**{email}** | **{phone}** | [{linkedin}](https://{linkedin}) | [{github}](https://{github})

---

## PROFESSIONAL SUMMARY
{summary}

---

## TECHNICAL SKILLS
**Core Competencies:** {skills_str}

---

## EXPERIENCE
**{job_title}** | {company}
*{job_dates}*
{job_desc}

---

## EDUCATION
**{degree}** | {university}
*{edu_dates}*
"""

st.markdown("### Resume Preview")
st.markdown("<div style='border: 1px solid #ddd; padding: 20px; border-radius: 5px; background-color: #1E2127;'>", unsafe_allow_html=True)
st.markdown(resume_md)
st.markdown("</div>", unsafe_allow_html=True)

# Generate DOCX
def create_docx(full_name, email, phone, linkedin, github, summary, skills_str, job_title, company, job_dates, job_desc, degree, university, edu_dates):
    doc = Document()
    
    # Name
    name_para = doc.add_paragraph()
    name_run = name_para.add_run(full_name)
    name_run.bold = True
    name_run.font.size = Pt(24)
    name_para.alignment = 1 # Center
    
    # Contact
    contact_para = doc.add_paragraph()
    contact_para.add_run(f"{email} | {phone} | {linkedin} | {github}")
    contact_para.alignment = 1
    
    # Summary
    doc.add_heading("PROFESSIONAL SUMMARY", level=1)
    doc.add_paragraph(summary)
    
    # Skills
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph(f"Core Competencies: {skills_str}")
    
    # Experience
    doc.add_heading("EXPERIENCE", level=1)
    exp_para = doc.add_paragraph()
    exp_para.add_run(f"{job_title} | {company}").bold = True
    exp_para.add_run(f"\n{job_dates}").italic = True
    for point in job_desc.split('\n'):
        if point.strip():
            # Remove leading hyphens for the bullet points since Word handles bullets
            clean_point = point.strip()
            if clean_point.startswith('- '):
                clean_point = clean_point[2:]
            doc.add_paragraph(clean_point, style='List Bullet')
            
    # Education
    doc.add_heading("EDUCATION", level=1)
    edu_para = doc.add_paragraph()
    edu_para.add_run(f"{degree} | {university}").bold = True
    edu_para.add_run(f"\n{edu_dates}").italic = True
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

docx_data = create_docx(full_name, email, phone, linkedin, github, summary, skills_str, job_title, company, job_dates, job_desc, degree, university, edu_dates)

st.markdown("---")
st.download_button(
    label="⬇️ Download as Word Document (.docx)",
    data=docx_data,
    file_name=f"{full_name.replace(' ', '_')}_Resume.docx",
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    use_container_width=True,
    type="primary"
)
